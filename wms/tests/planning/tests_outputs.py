from datetime import date
from pathlib import Path
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook

from wms.models import (
    CommunicationChannel,
    CommunicationDraft,
    CommunicationTemplate,
    PlanningArtifact,
    PlanningAssignment,
    PlanningAssignmentSource,
    PlanningCommunicationArtifact,
    PlanningFlightSnapshot,
    PlanningRun,
    PlanningShipmentSnapshot,
    PlanningVersion,
    PlanningVersionStatus,
    PlanningVolunteerSnapshot,
)
from wms.planning.communications import generate_version_drafts
from wms.planning.exports import export_version_pdf, export_version_workbook
from wms.planning.legacy_communications import CommunicationFamily
from wms.planning.stats import build_version_stats


class PlanningOutputTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="planner@example.com",
            email="planner@example.com",
            password="pass1234",  # pragma: allowlist secret
        )
        self.run = PlanningRun.objects.create(
            week_start="2026-03-09",
            week_end="2026-03-15",
            created_by=self.user,
        )
        self.shipment_snapshot = PlanningShipmentSnapshot.objects.create(
            run=self.run,
            shipment_reference="SHP-001",
            shipper_name="Hopital Saint Joseph",
            destination_iata="NSI",
            carton_count=4,
            equivalent_units=4,
            payload={
                "destination_city": "YAOUNDE",
                "legacy_type": "MM",
                "legacy_destinataire": "Centre Medical",
                "legacy_date_depart_mag": "2026-03-08",
            },
        )
        self.volunteer_snapshot = PlanningVolunteerSnapshot.objects.create(
            run=self.run,
            volunteer_label="Alice",
        )
        self.flight_snapshot = PlanningFlightSnapshot.objects.create(
            run=self.run,
            flight_number="AF123",
            departure_date="2026-03-10",
            destination_iata="CDG",
            payload={"departure_time": "11:10", "routing": "CDG-NSI"},
        )
        self.template = CommunicationTemplate.objects.create(
            label="Mail planning",
            channel=CommunicationChannel.EMAIL,
            scope=CommunicationFamily.EMAIL_ASF,
            subject="Planning v{{ version_number }} pour ASF interne",
            body="Bonjour ASF, semaine {{ week }}.",
        )

    def make_published_version(
        self,
        *,
        based_on=None,
        volunteer_label="Alice",
        flight_number="AF123",
        cartons=4,
    ):
        volunteer_snapshot = self.volunteer_snapshot
        if volunteer_label != self.volunteer_snapshot.volunteer_label:
            volunteer_snapshot = PlanningVolunteerSnapshot.objects.create(
                run=self.run,
                volunteer_label=volunteer_label,
            )

        flight_snapshot = self.flight_snapshot
        if flight_number != self.flight_snapshot.flight_number:
            flight_snapshot = PlanningFlightSnapshot.objects.create(
                run=self.run,
                flight_number=flight_number,
                departure_date="2026-03-11",
                destination_iata="NCE",
            )

        version = PlanningVersion.objects.create(
            run=self.run,
            status=PlanningVersionStatus.PUBLISHED,
            based_on=based_on,
            created_by=self.user,
        )
        PlanningAssignment.objects.create(
            version=version,
            shipment_snapshot=self.shipment_snapshot,
            volunteer_snapshot=volunteer_snapshot,
            flight_snapshot=flight_snapshot,
            assigned_carton_count=cartons,
            source=PlanningAssignmentSource.MANUAL,
            sequence=1,
        )
        return version

    def test_generate_drafts_and_strict_workbook_artifact_for_version(self):
        version = self.make_published_version()

        drafts = generate_version_drafts(version)
        artifact = export_version_workbook(version)
        stats = build_version_stats(version)

        self.assertEqual(len(drafts), 6)
        draft = CommunicationDraft.objects.get(
            version=version,
            family=CommunicationFamily.EMAIL_ASF,
        )
        self.assertEqual(draft.subject, "Planning v1 pour ASF interne v1")
        self.assertEqual(draft.body, "Bonjour ASF, semaine 11.")

        self.assertIsInstance(artifact, PlanningArtifact)
        self.assertEqual(artifact.artifact_type, "planning_workbook")
        self.assertTrue(artifact.file_path.endswith(".xlsx"))
        self.assertTrue(Path(artifact.file_path).exists())
        workbook = load_workbook(artifact.file_path)
        try:
            sheet = workbook.worksheets[0]
            week_anchor = sheet["A1"].value
            if hasattr(week_anchor, "date"):
                week_anchor = week_anchor.date()
            self.assertEqual(week_anchor, date(2026, 3, 9))
            self.assertEqual(sheet["Q1"].value, version.number)
            self.assertEqual(sheet["D35"].value, "Alice")
            self.assertEqual(sheet["F35"].value, "YAOUNDE")
            self.assertEqual(sheet["G35"].value, "NSI")
            self.assertEqual(sheet["H35"].value, "CDG-NSI")
            self.assertEqual(sheet["I35"].value, "AF 123")
            self.assertEqual(sheet["J35"].value, "11h10")
            self.assertEqual(sheet["K35"].value, "000001")
            self.assertEqual(sheet["L35"].value, 4)
            self.assertEqual(sheet["M35"].value, "MM")
            self.assertEqual(sheet["O35"].value, "08/03/26")
            self.assertEqual(sheet["P35"].value, "Hopital Saint Joseph")
            self.assertEqual(sheet["Q35"].value, "Centre Medical")
            self.assertTrue(sheet.row_dimensions[38].hidden)
        finally:
            workbook.close()

        self.assertEqual(stats["assignment_count"], 1)
        self.assertEqual(stats["carton_total"], 4)
        self.assertEqual(stats["volunteer_count"], 1)
        self.assertEqual(stats["flight_count"], 1)
        self.assertEqual(stats["manual_adjustment_count"], 1)
        self.assertEqual(stats["unassigned_count"], 0)
        self.assertEqual(
            stats["destination_breakdown"],
            [
                {
                    "destination_iata": "NSI",
                    "assignment_count": 1,
                    "carton_total": 4,
                    "equivalent_total": 4,
                }
            ],
        )
        self.assertEqual(
            stats["volunteer_breakdown"],
            [
                {
                    "volunteer_label": "Alice",
                    "assignment_count": 1,
                    "carton_total": 4,
                    "equivalent_total": 4,
                }
            ],
        )
        self.assertEqual(
            stats["flight_load_breakdown"],
            [
                {
                    "flight_snapshot_id": self.flight_snapshot.pk,
                    "flight_number": "AF123",
                    "departure_date": date(2026, 3, 10),
                    "departure_time": "11:10",
                    "destination_iata": "CDG",
                    "capacity_units": None,
                    "assignment_count": 1,
                    "carton_total": 4,
                    "equivalent_total": 4,
                    "remaining_units": None,
                    "utilization_pct": None,
                    "load_state": "unknown",
                    "load_state_label": "A renseigner",
                }
            ],
        )

    @mock.patch("wms.planning.exports.load_workbook")
    def test_export_version_workbook_closes_workbook_after_save(self, load_workbook_mock):
        version = self.make_published_version()
        workbook = mock.MagicMock()
        worksheet = mock.MagicMock()
        worksheet.max_row = 219
        worksheet.iter_rows.return_value = []
        workbook.worksheets = [worksheet]
        workbook.sheetnames = ["Planning SXX"]
        workbook.__getitem__.return_value = worksheet
        load_workbook_mock.return_value = workbook

        export_version_workbook(version)

        workbook.save.assert_called_once()
        workbook.close.assert_called_once_with()

    @mock.patch("wms.planning.exports.convert_workbook_to_pdf")
    def test_export_version_pdf_creates_pdf_artifact(self, convert_workbook_to_pdf_mock):
        version = self.make_published_version()

        def _fake_convert(workbook_path, pdf_path=None, *, strict=True):
            pdf_output = Path(pdf_path or Path(workbook_path).with_suffix(".pdf"))
            pdf_output.write_bytes(b"%PDF-1.4\n%")
            return pdf_output

        convert_workbook_to_pdf_mock.side_effect = _fake_convert

        artifact = export_version_pdf(version)

        self.assertIsInstance(artifact, PlanningArtifact)
        self.assertEqual(artifact.artifact_type, "planning_pdf")
        self.assertTrue(artifact.file_path.endswith(".pdf"))
        self.assertTrue(Path(artifact.file_path).exists())

    @mock.patch("wms.planning.exports.convert_workbook_to_pdf")
    def test_planning_export_records_pdf_artifact_health(self, convert_workbook_to_pdf_mock):
        version = self.make_published_version()

        def _fake_convert(workbook_path, pdf_path=None, *, strict=True):
            pdf_output = Path(pdf_path or Path(workbook_path).with_suffix(".pdf"))
            pdf_output.write_bytes(b"%PDF-1.4\n%")
            return pdf_output

        convert_workbook_to_pdf_mock.side_effect = _fake_convert

        artifact = export_version_pdf(version)

        self.assertEqual(artifact.artifact_type, "planning_pdf")
        health = PlanningCommunicationArtifact.objects.filter(
            planning_version=version,
            output_type="planning_pdf",
        ).latest("generated_at")
        self.assertEqual(health.status, "ready")
        self.assertEqual(health.output_type, "planning_pdf")
        self.assertTrue(health.file_name.endswith(".pdf"))

    def test_generate_drafts_aggregates_multiple_assignments_for_same_recipient(self):
        second_shipment = PlanningShipmentSnapshot.objects.create(
            run=self.run,
            shipment_reference="SHP-002",
            carton_count=2,
            equivalent_units=2,
        )
        second_flight = PlanningFlightSnapshot.objects.create(
            run=self.run,
            flight_number="AF456",
            departure_date="2026-03-11",
            destination_iata="NCE",
        )
        version = self.make_published_version()
        PlanningAssignment.objects.create(
            version=version,
            shipment_snapshot=second_shipment,
            volunteer_snapshot=self.volunteer_snapshot,
            flight_snapshot=second_flight,
            assigned_carton_count=2,
            source=PlanningAssignmentSource.MANUAL,
            sequence=2,
        )

        drafts = generate_version_drafts(version)

        self.assertEqual(len(drafts), 6)
        whatsapp_draft = next(
            draft for draft in drafts if draft.family == CommunicationFamily.WHATSAPP_BENEVOLE
        )
        self.assertEqual(whatsapp_draft.recipient_label, "Alice")
        self.assertIn("AF 123", whatsapp_draft.body)
        self.assertIn("AF 456", whatsapp_draft.body)
        self.assertIn("BE 000001", whatsapp_draft.body)
        self.assertIn("BE 000002", whatsapp_draft.body)

    def test_generate_drafts_creates_cancellation_message_for_removed_recipient(self):
        self.template.is_active = False
        self.template.save(update_fields=["is_active", "updated_at"])
        version_1 = self.make_published_version()
        version_2 = PlanningVersion.objects.create(
            run=self.run,
            status=PlanningVersionStatus.PUBLISHED,
            based_on=version_1,
            created_by=self.user,
        )

        drafts_v2 = generate_version_drafts(version_2)

        self.assertEqual(len(drafts_v2), 6)
        whatsapp_draft = next(
            draft for draft in drafts_v2 if draft.family == CommunicationFamily.WHATSAPP_BENEVOLE
        )
        self.assertEqual(whatsapp_draft.recipient_label, "Alice")
        self.assertIn("AF 123", whatsapp_draft.body)

    def test_generate_drafts_keeps_multiple_active_templates_on_same_channel(self):
        second_template = CommunicationTemplate.objects.create(
            label="Mail air france",
            channel=CommunicationChannel.EMAIL,
            scope=CommunicationFamily.EMAIL_AIRFRANCE,
            subject="Second {{ recipient_label }}",
            body="Second template pour {{ recipient_label }}",
        )
        version = self.make_published_version()

        drafts = generate_version_drafts(version)

        self.assertEqual(len(drafts), 6)
        self.assertEqual(
            sorted(draft.template_id for draft in drafts if draft.template_id),
            sorted([self.template.pk, second_template.pk]),
        )
        self.assertEqual(
            sorted(
                CommunicationDraft.objects.filter(
                    version=version, template__isnull=False
                ).values_list(
                    "template__label",
                    flat=True,
                )
            ),
            ["Mail air france", "Mail planning"],
        )

    def test_generate_drafts_keeps_separate_series_per_version(self):
        version_1 = self.make_published_version()
        version_2 = self.make_published_version(
            based_on=version_1,
            volunteer_label="Bob",
            flight_number="AF456",
            cartons=5,
        )

        drafts_v1 = generate_version_drafts(version_1)
        drafts_v2 = generate_version_drafts(version_2)

        self.assertEqual(len(drafts_v1), 6)
        self.assertEqual(len(drafts_v2), 7)
        self.assertEqual(
            sorted(
                CommunicationDraft.objects.filter(version=version_2).values_list(
                    "family",
                    "recipient_label",
                )
            ),
            sorted(
                [
                    (CommunicationFamily.EMAIL_AIRFRANCE, "Air France"),
                    (CommunicationFamily.EMAIL_ASF, "ASF interne"),
                    (CommunicationFamily.EMAIL_CORRESPONDANT, "YAOUNDE"),
                    (CommunicationFamily.EMAIL_DESTINATAIRE, "Centre Medical"),
                    (CommunicationFamily.EMAIL_EXPEDITEUR, "Hopital Saint Joseph"),
                    (CommunicationFamily.WHATSAPP_BENEVOLE, "Alice"),
                    (CommunicationFamily.WHATSAPP_BENEVOLE, "Bob"),
                ]
            ),
        )
