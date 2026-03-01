from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.core.files.base import ContentFile
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    GeneratedPrintArtifact,
    GeneratedPrintArtifactItem,
    GeneratedPrintArtifactStatus,
    PrintPack,
)
from .print_pack_excel import fill_workbook_cells
from .print_pack_graph import convert_excel_to_pdf_via_graph
from .print_pack_pdf import merge_pdf_documents


class PrintPackEngineError(RuntimeError):
    """Raised when pack generation cannot be completed."""


@dataclass(frozen=True)
class PackXlsxDocument:
    filename: str
    payload: bytes


def _clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _join_non_empty(*values, separator=" "):
    parts = [_clean_text(value) for value in values if _clean_text(value)]
    return separator.join(parts)


def _first_line(value):
    text = _clean_text(value)
    if not text:
        return ""
    return text.splitlines()[0].strip()


def _unique_non_empty(values):
    result = []
    seen = set()
    for value in values:
        text = _clean_text(value)
        if not text:
            continue
        lowered = text.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        result.append(text)
    return result


def _format_weight_label(total_weight_g):
    if total_weight_g is None:
        total_weight_g = 0
    weight_kg = float(total_weight_g) / 1000.0
    if weight_kg.is_integer():
        return f"{int(weight_kg)} kg"
    return f"{weight_kg:.2f}".rstrip("0").rstrip(".") + " kg"


def _build_contact_payload(*, contact, fallback_name, default_country=""):
    base_name = _clean_text(fallback_name)
    payload = {
        "full_name": base_name,
        "title_name": base_name,
        "structure_name": "",
        "postal_address": "",
        "postal_code": "",
        "city": "",
        "country": _clean_text(default_country),
        "phone_1": "",
        "phone_2": "",
        "phone_3": "",
        "email_1": "",
        "email_2": "",
        "email_3": "",
        "emergency_contact": "",
        "postal_address_full": "",
        "contact_primary": "",
    }
    if contact is None:
        return payload

    first_name = _clean_text(getattr(contact, "first_name", ""))
    last_name = _clean_text(getattr(contact, "last_name", ""))
    title = _clean_text(getattr(contact, "title", ""))
    if first_name or last_name:
        payload["title_name"] = _join_non_empty(
            title,
            first_name,
            last_name.upper() if last_name else "",
        )
    elif _clean_text(getattr(contact, "name", "")):
        payload["title_name"] = _clean_text(getattr(contact, "name", ""))
    if not payload["title_name"]:
        payload["title_name"] = base_name
    payload["full_name"] = payload["title_name"]

    organization = getattr(contact, "organization", None)
    contact_name = _clean_text(getattr(contact, "name", ""))
    if organization is not None and _clean_text(getattr(organization, "name", "")):
        payload["structure_name"] = _clean_text(getattr(organization, "name", ""))
    else:
        contact_type = _clean_text(getattr(contact, "contact_type", ""))
        if contact_type == "organization":
            payload["structure_name"] = contact_name

    address = None
    if hasattr(contact, "get_effective_address"):
        address = contact.get_effective_address()
    if address is not None:
        payload["postal_address"] = _join_non_empty(
            getattr(address, "address_line1", ""),
            getattr(address, "address_line2", ""),
            separator=", ",
        )
        payload["postal_code"] = _clean_text(getattr(address, "postal_code", ""))
        payload["city"] = _clean_text(getattr(address, "city", ""))
        if _clean_text(getattr(address, "country", "")):
            payload["country"] = _clean_text(getattr(address, "country", ""))

    phones = _unique_non_empty(
        [
            getattr(contact, "phone", ""),
            getattr(contact, "phone2", ""),
            getattr(address, "phone", "") if address is not None else "",
        ]
    )
    emails = _unique_non_empty(
        [
            getattr(contact, "email", ""),
            getattr(contact, "email2", ""),
            getattr(address, "email", "") if address is not None else "",
        ]
    )
    payload["phone_1"] = phones[0] if len(phones) > 0 else ""
    payload["phone_2"] = phones[1] if len(phones) > 1 else ""
    payload["phone_3"] = phones[2] if len(phones) > 2 else ""
    payload["email_1"] = emails[0] if len(emails) > 0 else ""
    payload["email_2"] = emails[1] if len(emails) > 1 else ""
    payload["email_3"] = emails[2] if len(emails) > 2 else ""
    payload["emergency_contact"] = _first_line(getattr(contact, "notes", ""))
    if not payload["emergency_contact"]:
        payload["emergency_contact"] = _clean_text(getattr(contact, "role", ""))

    city_line = _join_non_empty(payload["postal_code"], payload["city"])
    payload["postal_address_full"] = _join_non_empty(
        payload["postal_address"],
        city_line,
        payload["country"],
        separator=", ",
    )
    payload["contact_primary"] = _join_non_empty(
        payload["phone_1"],
        payload["email_1"],
        separator=", ",
    )
    return payload


def _resolve_root_category_name(product):
    category = getattr(product, "category", None)
    if category is None:
        return ""
    visited_ids = set()
    while getattr(category, "parent", None) is not None:
        category_id = getattr(category, "id", None)
        if category_id in visited_ids:
            break
        if category_id is not None:
            visited_ids.add(category_id)
        category = category.parent
    return getattr(category, "name", "") or ""


def _build_recipient_payload(shipment):
    return _build_contact_payload(
        contact=getattr(shipment, "recipient_contact_ref", None),
        fallback_name=getattr(shipment, "recipient_name", ""),
        default_country=getattr(shipment, "destination_country", ""),
    )


def _build_mapping_payload(*, shipment=None, carton=None, document=None):
    generated_on = timezone.localdate()
    payload = {
        "shipment": {},
        "carton": {},
        "document": {
            "doc_type": getattr(document, "doc_type", ""),
            "variant": getattr(document, "variant", ""),
            "generated_on": generated_on,
        },
    }
    if shipment is not None:
        origin_city = "PARIS"
        origin_iata = "CDG"
        carton_count_attr = getattr(shipment, "carton_count", None)
        if carton_count_attr is not None:
            carton_total_count = carton_count_attr
        elif hasattr(shipment, "carton_set"):
            carton_total_count = shipment.carton_set.count()
        else:
            carton_total_count = 1
        destination = getattr(shipment, "destination", None)
        destination_iata = ""
        destination_city = ""
        if destination is not None:
            destination_iata = getattr(destination, "iata_code", "") or ""
            destination_city = getattr(destination, "city", "") or ""

        shipment_items = []
        total_weight_g = 0
        hors_format_total_count = 0
        if hasattr(shipment, "carton_set"):
            carton_qs = shipment.carton_set.all().order_by("code")
            for carton_position, shipment_carton in enumerate(carton_qs, start=1):
                carton_items = shipment_carton.cartonitem_set.select_related(
                    "product_lot__product__category__parent",
                    "product_lot__location",
                )
                for carton_item in carton_items:
                    product = carton_item.product_lot.product
                    root_category = _resolve_root_category_name(product)
                    quantity = carton_item.quantity
                    product_weight = getattr(product, "weight_g", None) or 0
                    total_weight_g += product_weight * quantity
                    if root_category.upper() == "HF":
                        hors_format_total_count += quantity
                    shipment_items.append(
                        {
                            "carton_code": shipment_carton.code,
                            "carton_position": carton_position,
                            "category_root": root_category,
                            "brand": product.brand or "",
                            "product_name": product.name,
                            "quantity": quantity,
                            "expires_on": carton_item.product_lot.expires_on,
                        }
                    )
        payload["shipment"] = {
            "id": shipment.id,
            "reference": shipment.reference,
            "carton_total_count": carton_total_count,
            "origin_city": origin_city,
            "origin_iata": origin_iata,
            "destination_iata": destination_iata,
            "destination_city": destination_city,
            "total_weight_g": total_weight_g,
            "total_weight_label": _format_weight_label(total_weight_g),
            "hors_format_total_count": hors_format_total_count,
            "shipper_name": shipment.shipper_name,
            "shipper": _build_contact_payload(
                contact=getattr(shipment, "shipper_contact_ref", None),
                fallback_name=shipment.shipper_name,
            ),
            "recipient_name": shipment.recipient_name,
            "recipient": _build_recipient_payload(shipment),
            "correspondent_name": shipment.correspondent_name,
            "correspondent": _build_contact_payload(
                contact=getattr(shipment, "correspondent_contact_ref", None),
                fallback_name=shipment.correspondent_name,
                default_country=shipment.destination_country,
            ),
            "destination_address": shipment.destination_address,
            "destination_country": shipment.destination_country,
            "requested_delivery_date": shipment.requested_delivery_date,
            "notes": shipment.notes,
            "items": shipment_items,
        }
    if carton is not None:
        carton_position = ""
        shipment_context = shipment or getattr(carton, "shipment", None)
        if shipment_context is not None and hasattr(shipment_context, "carton_set"):
            cartons = list(shipment_context.carton_set.all().order_by("code"))
            for index, shipment_carton in enumerate(cartons, start=1):
                if shipment_carton.id == carton.id:
                    carton_position = index
                    break
        items = []
        if hasattr(carton, "cartonitem_set"):
            for carton_item in carton.cartonitem_set.select_related(
                "product_lot__product",
                "product_lot__location",
            ):
                product = carton_item.product_lot.product
                location = carton_item.product_lot.location
                if location:
                    location_label = f"{location.zone} - {location.aisle} - {location.shelf}"
                else:
                    location_label = ""
                items.append(
                    {
                        "category_root": _resolve_root_category_name(product),
                        "product_name": product.name,
                        "brand": product.brand or "",
                        "quantity": carton_item.quantity,
                        "expires_on": carton_item.product_lot.expires_on,
                        "location": location_label,
                    }
                )
        payload["carton"] = {
            "id": carton.id,
            "code": carton.code,
            "position": carton_position,
            "items": items,
        }
    return payload


def _resolve_template_search_dirs():
    configured_dirs = getattr(settings, "PRINT_PACK_TEMPLATE_DIRS", None)
    if configured_dirs is None:
        configured_dirs = [str(Path(settings.BASE_DIR) / "data" / "print_templates")]
    if isinstance(configured_dirs, (str, Path)):
        configured_dirs = [configured_dirs]

    search_dirs = []
    for entry in configured_dirs:
        value = str(entry or "").strip()
        if not value:
            continue
        search_dirs.append(Path(value))
    return search_dirs


def _resolve_template_filename(document):
    pack_code = _clean_text(getattr(getattr(document, "pack", None), "code", ""))
    doc_type = _clean_text(getattr(document, "doc_type", ""))
    variant = _clean_text(getattr(document, "variant", ""))
    if not (pack_code and doc_type and variant):
        return ""
    return f"{pack_code}__{doc_type}__{variant}.xlsx"


def _read_template_bytes_from_search_dirs(document):
    template_name = _resolve_template_filename(document)
    if not template_name:
        return None, []

    attempted_paths = []
    for directory in _resolve_template_search_dirs():
        candidate = directory / template_name
        attempted_paths.append(str(candidate))
        if candidate.exists() and candidate.is_file():
            with candidate.open("rb") as stream:
                return stream.read(), attempted_paths
    return None, attempted_paths


def _render_document_xlsx_bytes(*, document, shipment=None, carton=None):
    template_bytes = None
    if document.xlsx_template_file:
        with document.xlsx_template_file.open("rb") as stream:
            template_bytes = stream.read()
    else:
        template_bytes, attempted_paths = _read_template_bytes_from_search_dirs(document)
        if template_bytes is None:
            attempts_text = ", ".join(attempted_paths) if attempted_paths else "none"
            raise PrintPackEngineError(
                "Missing xlsx template file for doc_type="
                f"{document.doc_type}. Searched: {attempts_text}"
            )

    workbook = load_workbook(BytesIO(template_bytes))
    mappings = list(document.cell_mappings.order_by("sequence", "id"))
    payload = _build_mapping_payload(
        shipment=shipment,
        carton=carton,
        document=document,
    )
    fill_workbook_cells(workbook, mappings, payload)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _artifact_basename(*, pack_code):
    stamp = timezone.now().strftime("%Y%m%d%H%M%S")
    return f"print-pack-{pack_code}-{stamp}"


def _document_render_targets(*, document, shipment=None, carton=None):
    if (
        document.doc_type == "destination_label"
        and (document.variant or "") == "all_labels"
        and shipment is not None
        and hasattr(shipment, "carton_set")
    ):
        cartons = list(shipment.carton_set.all().order_by("code"))
        if cartons:
            return [(shipment, shipment_carton) for shipment_carton in cartons]
    return [(shipment, carton)]


def _resolve_pack_and_documents(*, pack_code, variant=None):
    pack = PrintPack.objects.filter(code=pack_code, active=True).first()
    if pack is None:
        raise PrintPackEngineError(f"Unknown active pack: {pack_code}")

    documents_qs = pack.documents.filter(enabled=True)
    if variant:
        documents_qs = documents_qs.filter(variant=variant)
    documents = list(documents_qs.order_by("sequence", "id"))
    if not documents:
        raise PrintPackEngineError(f"No enabled documents configured for pack {pack_code}.")
    return pack, documents


def render_pack_xlsx_documents(*, pack_code, shipment=None, carton=None, variant=None):
    pack, documents = _resolve_pack_and_documents(pack_code=pack_code, variant=variant)

    xlsx_documents = []
    for document in documents:
        targets = _document_render_targets(
            document=document,
            shipment=shipment,
            carton=carton,
        )
        for target_index, (target_shipment, target_carton) in enumerate(targets, start=1):
            xlsx_bytes = _render_document_xlsx_bytes(
                document=document,
                shipment=target_shipment,
                carton=target_carton,
            )
            filename_suffix = f"-{target_index}" if len(targets) > 1 else ""
            xlsx_name = (
                f"{pack.code}-{document.doc_type}-{document.id}{filename_suffix}.xlsx"
            )
            xlsx_documents.append(PackXlsxDocument(filename=xlsx_name, payload=xlsx_bytes))
    return xlsx_documents


def generate_pack(*, pack_code, shipment=None, carton=None, user=None, variant=None):
    pack, documents = _resolve_pack_and_documents(pack_code=pack_code, variant=variant)

    artifact = GeneratedPrintArtifact.objects.create(
        shipment=shipment,
        carton=carton,
        pack_code=pack.code,
        status=GeneratedPrintArtifactStatus.GENERATED,
        created_by=user,
    )

    generated_pdfs = []
    for document in documents:
        targets = _document_render_targets(
            document=document,
            shipment=shipment,
            carton=carton,
        )
        for target_index, (target_shipment, target_carton) in enumerate(targets, start=1):
            xlsx_bytes = _render_document_xlsx_bytes(
                document=document,
                shipment=target_shipment,
                carton=target_carton,
            )
            filename_suffix = f"-{target_index}" if len(targets) > 1 else ""
            xlsx_name = (
                f"{pack.code}-{document.doc_type}-{document.id}{filename_suffix}.xlsx"
            )
            pdf_name = f"{pack.code}-{document.doc_type}-{document.id}{filename_suffix}.pdf"
            pdf_bytes = convert_excel_to_pdf_via_graph(
                xlsx_bytes=xlsx_bytes,
                filename=xlsx_name,
            )
            generated_pdfs.append(pdf_bytes)

            item = GeneratedPrintArtifactItem.objects.create(
                artifact=artifact,
                doc_type=document.doc_type,
                variant=document.variant or "",
                sequence=document.sequence,
            )
            item.source_xlsx_file.save(
                xlsx_name,
                ContentFile(xlsx_bytes),
                save=False,
            )
            item.generated_pdf_file.save(
                pdf_name,
                ContentFile(pdf_bytes),
                save=False,
            )
            item.save(update_fields=["source_xlsx_file", "generated_pdf_file"])

    if len(generated_pdfs) == 1:
        merged_pdf = generated_pdfs[0]
    else:
        merged_pdf = merge_pdf_documents(generated_pdfs)

    artifact_name = f"{_artifact_basename(pack_code=pack.code)}.pdf"
    artifact.pdf_file.save(artifact_name, ContentFile(merged_pdf), save=False)
    artifact.status = GeneratedPrintArtifactStatus.SYNC_PENDING
    artifact.save(update_fields=["pdf_file", "status"])
    return artifact
