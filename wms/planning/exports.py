from __future__ import annotations

import os
import tempfile
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from tools.planning_comm_helper import excel_pdf
from tools.planning_comm_helper.planning_pdf import (
    PlanningPdfConversionError,
    convert_workbook_to_pdf,
)
from wms.models import PlanningArtifact, PlanningVersion
from wms.planning.artifact_health import (
    PLANNING_ARTIFACT_STATUS_FAILED,
    PLANNING_ARTIFACT_STATUS_READY,
    PLANNING_WORKBOOK_BACKEND,
    artifact_file_name,
    record_planning_artifact_result,
)
from wms.planning.legacy_communications import (
    format_be_number,
    format_heure_hh_mm,
    format_vol_display,
)

PLANNING_WORKBOOK_ARTIFACT = "planning_workbook"
PLANNING_PDF_ARTIFACT = "planning_pdf"
PLANNING_TEMPLATE_FILENAME = "Planning-maquette.xlsx"

_PLAN_DAY_BLOCKS: dict[int, tuple[int, int]] = {
    0: (4, 32),
    1: (35, 63),
    2: (66, 94),
    3: (97, 125),
    4: (128, 156),
    5: (159, 187),
    6: (190, 218),
}

_PLAN_KEEP_ROWS: set[int] = {
    3,
    4,
    5,
    6,
    31,
    32,
    33,
    34,
    35,
    36,
    37,
    62,
    63,
    64,
    65,
    66,
    67,
    68,
    93,
    94,
    95,
    96,
    97,
    98,
    99,
    124,
    125,
    126,
    127,
    128,
    129,
    130,
    155,
    156,
    157,
    158,
    159,
    160,
    161,
    186,
    187,
    188,
    189,
    190,
    191,
    192,
    217,
    218,
    219,
}

_PLAN_MIDDLE_MOVES: tuple[tuple[int, int, int], ...] = (
    (17, 4, 32),
    (48, 35, 63),
    (79, 66, 94),
    (110, 97, 125),
    (141, 128, 156),
    (172, 159, 187),
    (203, 190, 218),
)

_EXPORT_PDF_ERRORS = (
    PlanningPdfConversionError,
    FileNotFoundError,
    OSError,
    PermissionError,
    RuntimeError,
    ValueError,
    TypeError,
)


class PlanningExportError(RuntimeError):
    """Raised when a strict planning export cannot be generated."""


@dataclass(frozen=True)
class _PlanningExportRow:
    day_idx: int
    departure_date: date | None
    departure_time_sort_key: tuple[int, int]
    flight_group_key: tuple[date | None, tuple[int, int], str]
    volunteer_display: str
    city: str
    iata: str
    routing: str
    flight_display: str
    departure_time_display: str
    be_number: str
    carton_count: int
    shipment_type: str
    departure_mag_display: str
    shipper_name: str
    recipient_name: str


def _planning_output_dir() -> Path:
    base_dir = Path(os.getenv("ASF_TMP_DIR") or tempfile.gettempdir())
    output_dir = base_dir / "asf_wms_planning_exports"
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[2]


def _planning_template_path() -> Path:
    template_path = _repo_root() / "data" / "planning_templates" / PLANNING_TEMPLATE_FILENAME
    if not template_path.exists():
        raise PlanningExportError(f"Planning template not found: {template_path}")
    return template_path


def _planning_basename(version: PlanningVersion) -> str:
    version_token = f"v{version.number}" if version.number else f"id{version.pk}"
    return f"planning-run-{version.run_id}-{version_token}"


def _safe_int(value) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _normalize_text(value: object, *, uppercase: bool = False) -> str:
    text = str(value or "").strip()
    return text.upper() if uppercase else text


def _coerce_departure_time_sort_key(value: object) -> tuple[int, int]:
    text = str(value or "").strip().lower().replace("h", ":")
    if not text:
        return (99, 59)
    hour, _, minute = text.partition(":")
    try:
        return (int(hour), int((minute or "0")[:2]))
    except ValueError:
        return (99, 59)


def _format_departure_mag(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%y")
    text = str(value or "").strip()
    if not text:
        return ""
    for parser in (date.fromisoformat,):
        try:
            return parser(text).strftime("%d/%m/%y")
        except (TypeError, ValueError):
            continue
    return text


def _coerce_date_value(value: object) -> date | object:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return value
    try:
        return date.fromisoformat(text)
    except ValueError:
        return value


def _build_export_rows(version: PlanningVersion) -> list[_PlanningExportRow]:
    assignments = list(
        version.assignments.select_related(
            "shipment_snapshot",
            "volunteer_snapshot",
            "flight_snapshot",
        ).order_by(
            "flight_snapshot__departure_date",
            "flight_snapshot__flight_number",
            "sequence",
            "id",
        )
    )
    rows: list[_PlanningExportRow] = []
    for assignment in assignments:
        shipment = assignment.shipment_snapshot
        volunteer = assignment.volunteer_snapshot
        flight = assignment.flight_snapshot
        shipment_payload = shipment.payload if shipment and shipment.payload else {}
        flight_payload = flight.payload if flight and flight.payload else {}

        departure_date = flight.departure_date if flight else None
        departure_time = flight_payload.get("departure_time") if flight else ""
        departure_time_sort_key = _coerce_departure_time_sort_key(departure_time)
        flight_display = format_vol_display(flight.flight_number if flight else "")
        city = _normalize_text(
            shipment_payload.get("destination_city")
            or (shipment.destination_iata if shipment else "")
            or (flight.destination_iata if flight else ""),
            uppercase=True,
        )
        iata = _normalize_text(
            shipment.destination_iata if shipment else (flight.destination_iata if flight else ""),
            uppercase=True,
        )
        routing = _normalize_text(flight_payload.get("routing"), uppercase=True).replace(",", "-")
        shipper_name = _normalize_text(
            shipment.shipper_name if shipment else shipment_payload.get("legacy_expediteur")
        )
        recipient_name = _normalize_text(shipment_payload.get("legacy_destinataire"))
        rows.append(
            _PlanningExportRow(
                day_idx=departure_date.weekday() if departure_date else 0,
                departure_date=departure_date,
                departure_time_sort_key=departure_time_sort_key,
                flight_group_key=(departure_date, departure_time_sort_key, flight_display),
                volunteer_display=_normalize_text(
                    volunteer.volunteer_label if volunteer else "",
                    uppercase=False,
                ),
                city=city,
                iata=iata,
                routing=routing,
                flight_display=flight_display,
                departure_time_display=format_heure_hh_mm(departure_time),
                be_number=format_be_number(shipment.shipment_reference if shipment else ""),
                carton_count=_safe_int(assignment.assigned_carton_count),
                shipment_type=_normalize_text(shipment_payload.get("legacy_type"), uppercase=True),
                departure_mag_display=_format_departure_mag(
                    shipment_payload.get("legacy_date_depart_mag")
                ),
                shipper_name=shipper_name,
                recipient_name=recipient_name,
            )
        )
    rows.sort(
        key=lambda row: (
            row.day_idx,
            row.departure_date or date.max,
            row.departure_time_sort_key,
            row.flight_display,
            row.city,
            row.be_number,
            row.volunteer_display.lower(),
        )
    )
    return rows


def _planning_sheet(workbook):
    for name in workbook.sheetnames:
        if name.lower().startswith("planning"):
            return workbook[name]
    return workbook.worksheets[0]


def _reset_planning_grid(ws_plan) -> None:
    for row in ws_plan.iter_rows(min_row=3, max_row=ws_plan.max_row, min_col=4, max_col=17):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if cell.coordinate in {"K219", "L219"}:
                continue
            cell.value = None
        ws_plan.row_dimensions[row[0].row].hidden = False


def _hide_non_keep_rows(ws_plan, *, start: int, end: int, keep_rows: set[int]) -> None:
    for row_idx in range(start, end + 1):
        if row_idx not in keep_rows:
            ws_plan.row_dimensions[row_idx].hidden = True


def _write_planning_row(
    ws_plan,
    *,
    row_idx: int,
    row: _PlanningExportRow,
    volunteer_value: str,
    is_first: bool,
) -> None:
    ws_plan.cell(row=row_idx, column=4).value = volunteer_value
    if is_first:
        ws_plan.cell(row=row_idx, column=6).value = row.city
        ws_plan.cell(row=row_idx, column=7).value = row.iata
        ws_plan.cell(row=row_idx, column=8).value = row.routing
        ws_plan.cell(row=row_idx, column=9).value = row.flight_display
        ws_plan.cell(row=row_idx, column=10).value = row.departure_time_display
    ws_plan.cell(row=row_idx, column=11).value = row.be_number
    ws_plan.cell(row=row_idx, column=12).value = row.carton_count
    ws_plan.cell(row=row_idx, column=13).value = row.shipment_type
    ws_plan.cell(row=row_idx, column=15).value = row.departure_mag_display
    ws_plan.cell(row=row_idx, column=16).value = row.shipper_name
    ws_plan.cell(row=row_idx, column=17).value = row.recipient_name


def _group_day_rows(
    rows: Iterable[_PlanningExportRow],
) -> dict[tuple[date | None, tuple[int, int], str], list[_PlanningExportRow]]:
    grouped: dict[tuple[date | None, tuple[int, int], str], list[_PlanningExportRow]] = defaultdict(
        list
    )
    for row in rows:
        grouped[row.flight_group_key].append(row)
    return grouped


def _populate_planning_sheet(ws_plan, *, rows: list[_PlanningExportRow]) -> None:
    rows_by_day: dict[int, list[_PlanningExportRow]] = defaultdict(list)
    for row in rows:
        rows_by_day[row.day_idx].append(row)

    for day_idx in range(7):
        block = _PLAN_DAY_BLOCKS.get(day_idx)
        if block is None:
            continue
        start, end = block
        current_row = start
        day_rows = rows_by_day.get(day_idx, [])
        if not day_rows:
            _hide_non_keep_rows(ws_plan, start=start, end=end, keep_rows=_PLAN_KEEP_ROWS)
            continue

        grouped = _group_day_rows(day_rows)
        ordered_keys = sorted(grouped.keys(), key=lambda key: (key[0] or date.max, key[1], key[2]))
        for key in ordered_keys:
            flight_rows = grouped[key]
            if current_row > end:
                break
            volunteer_values = list(
                dict.fromkeys(
                    value for value in (row.volunteer_display for row in flight_rows) if value
                )
            )
            for index, row in enumerate(flight_rows):
                if current_row > end:
                    break
                volunteer_value = volunteer_values[index] if index < len(volunteer_values) else ""
                _write_planning_row(
                    ws_plan,
                    row_idx=current_row,
                    row=row,
                    volunteer_value=volunteer_value,
                    is_first=index == 0,
                )
                current_row += 1
            current_row += 2

        if current_row <= end:
            _hide_non_keep_rows(ws_plan, start=current_row, end=end, keep_rows=_PLAN_KEEP_ROWS)


def _move_cell_value_to_visible_middle(
    ws_plan,
    *,
    src_row: int,
    start: int,
    end: int,
    col_letter: str = "A",
) -> None:
    source = ws_plan[f"{col_letter}{src_row}"]
    if source.value is None:
        return
    visible_rows = [
        row_idx
        for row_idx in range(start, end + 1)
        if not bool(ws_plan.row_dimensions[row_idx].hidden)
    ]
    if not visible_rows:
        return
    destination_row = visible_rows[len(visible_rows) // 2]
    if destination_row == src_row:
        return
    destination = ws_plan[f"{col_letter}{destination_row}"]
    destination.value = source.value
    destination._style = source._style
    source.value = None


def _apply_planning_layout(ws_plan) -> None:
    for col_letter in ("B", "C", "E", "G", "N"):
        ws_plan.column_dimensions[col_letter].hidden = True
    for col_letter in ("P", "Q"):
        column = ws_plan[col_letter]
        max_len = max((len(str(cell.value)) for cell in column if cell.value), default=10)
        ws_plan.column_dimensions[col_letter].width = max(10, min(max_len + 2, 40))
    for src_row, start, end in _PLAN_MIDDLE_MOVES:
        _move_cell_value_to_visible_middle(
            ws_plan,
            src_row=src_row,
            start=start,
            end=end,
        )


def _write_metadata(ws_plan, *, version: PlanningVersion) -> None:
    ws_plan["A1"] = _coerce_date_value(version.run.week_start)
    ws_plan["Q1"] = version.number


def _artifact_label(version: PlanningVersion, *, suffix: str) -> str:
    version_token = f"v{version.number}" if version.number else f"id{version.pk}"
    return f"Planning {suffix} {version_token}"


def _upsert_artifact(
    *,
    version: PlanningVersion,
    artifact_type: str,
    label: str,
    file_path: Path,
) -> PlanningArtifact:
    artifact = version.artifacts.filter(artifact_type=artifact_type).order_by("id").first()
    if artifact is None:
        return PlanningArtifact.objects.create(
            version=version,
            artifact_type=artifact_type,
            label=label,
            file_path=str(file_path),
        )
    artifact.label = label
    artifact.file_path = str(file_path)
    artifact.save(update_fields=["label", "file_path", "generated_at"])
    return artifact


def export_version_workbook(version: PlanningVersion) -> PlanningArtifact:
    template_path = _planning_template_path()
    output_path = _planning_output_dir() / f"{_planning_basename(version)}.xlsx"
    workbook = load_workbook(template_path)
    try:
        try:
            ws_plan = _planning_sheet(workbook)
            _reset_planning_grid(ws_plan)
            _write_metadata(ws_plan, version=version)
            _populate_planning_sheet(ws_plan, rows=_build_export_rows(version))
            _apply_planning_layout(ws_plan)
            workbook.save(output_path)
            artifact = _upsert_artifact(
                version=version,
                artifact_type=PLANNING_WORKBOOK_ARTIFACT,
                label=_artifact_label(version, suffix="XLSX"),
                file_path=output_path,
            )
            record_planning_artifact_result(
                version=version,
                output_type=PLANNING_WORKBOOK_ARTIFACT,
                status=PLANNING_ARTIFACT_STATUS_READY,
                backend=PLANNING_WORKBOOK_BACKEND,
                file_name=artifact_file_name(output_path),
                payload={"artifact_id": artifact.pk},
            )
            return artifact
        except Exception as exc:
            record_planning_artifact_result(
                version=version,
                output_type=PLANNING_WORKBOOK_ARTIFACT,
                status=PLANNING_ARTIFACT_STATUS_FAILED,
                backend=PLANNING_WORKBOOK_BACKEND,
                file_name=artifact_file_name(output_path),
                error_message=str(exc),
            )
            raise
    finally:
        workbook.close()


def _export_pdf_artifact(
    *,
    version: PlanningVersion,
    workbook_artifact: PlanningArtifact,
) -> PlanningArtifact:
    workbook_path = Path(workbook_artifact.file_path)
    pdf_path = _planning_output_dir() / f"{_planning_basename(version)}.pdf"
    try:
        generated = Path(convert_workbook_to_pdf(workbook_path, pdf_path, strict=True))
    except _EXPORT_PDF_ERRORS as exc:
        record_planning_artifact_result(
            version=version,
            output_type=PLANNING_PDF_ARTIFACT,
            status=PLANNING_ARTIFACT_STATUS_FAILED,
            backend=excel_pdf.pdf_backend_name(),
            file_name=artifact_file_name(pdf_path),
            error_message=str(exc),
            payload={"workbook_artifact_id": workbook_artifact.pk},
        )
        raise PlanningExportError("Planning PDF indisponible.") from exc
    if not generated.exists():
        record_planning_artifact_result(
            version=version,
            output_type=PLANNING_PDF_ARTIFACT,
            status=PLANNING_ARTIFACT_STATUS_FAILED,
            backend=excel_pdf.pdf_backend_name(),
            file_name=artifact_file_name(pdf_path),
            error_message="Generated PDF file is missing.",
            payload={"workbook_artifact_id": workbook_artifact.pk},
        )
        raise PlanningExportError("Planning PDF indisponible.")
    artifact = _upsert_artifact(
        version=version,
        artifact_type=PLANNING_PDF_ARTIFACT,
        label=_artifact_label(version, suffix="PDF"),
        file_path=generated,
    )
    record_planning_artifact_result(
        version=version,
        output_type=PLANNING_PDF_ARTIFACT,
        status=PLANNING_ARTIFACT_STATUS_READY,
        backend=excel_pdf.pdf_backend_name(),
        file_name=artifact_file_name(generated),
        payload={
            "artifact_id": artifact.pk,
            "workbook_artifact_id": workbook_artifact.pk,
        },
    )
    return artifact


def export_version_pdf(version: PlanningVersion) -> PlanningArtifact:
    workbook_artifact = export_version_workbook(version)
    return _export_pdf_artifact(version=version, workbook_artifact=workbook_artifact)


def export_version_artifacts(version: PlanningVersion) -> dict[str, PlanningArtifact]:
    workbook_artifact = export_version_workbook(version)
    pdf_artifact = _export_pdf_artifact(version=version, workbook_artifact=workbook_artifact)
    return {
        PLANNING_WORKBOOK_ARTIFACT: workbook_artifact,
        PLANNING_PDF_ARTIFACT: pdf_artifact,
    }
